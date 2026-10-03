"""Node types that reach outside the RPA graph: shell commands and spreadsheets.

Kept apart from :mod:`apps.engine.builtin_nodes` because both of these cross a trust
boundary — one executes a program, one opens a file a user picked — and both
deserve their guardrails stated in one place instead of scattered.

**``run_shell`` does not interpret its command through a shell by default.** It
takes an argv list, so a variable holding a filename with a space or a semicolon
cannot turn into an extra command. ``shell=True`` is available but has to be
asked for by name, which is the difference between a capability and an accident.

**``excel`` uses openpyxl, and says so when it is missing.** The alternative —
silently returning an empty sheet — is the failure mode that makes an automation
look like it worked.
"""

from __future__ import annotations

import shlex
import subprocess
import time
from pathlib import Path
from typing import Any

from .registry import BaseNode, NodeRegistry, NodeSpec, ParamSpec
from .schema import NodeError

#: Refuse to run anything that would outlive the flow by more than this unless
#: the author raised the node's own timeout.
DEFAULT_SHELL_TIMEOUT = 300


class RunShellNode(BaseNode):
    """Execute a program and capture its output.

    ``command`` accepts either a list of arguments (preferred) or a string that
    is split with :func:`shlex.split` — not handed to ``/bin/sh``. Set
    ``shell`` only when the command genuinely needs shell features such as
    pipes or globbing; that flag turns a filename into an attack surface.
    """

    def execute(self) -> dict[str, Any]:
        raw = self.param("command", None)
        if not raw:
            raise NodeError("run_shell 需要 command 参数")

        use_shell = bool(self.param("shell", False))
        if isinstance(raw, (list, tuple)):
            argv = [str(part) for part in raw]
        else:
            text = str(raw).strip()
            if not text:
                raise NodeError("command 为空")
            if use_shell:
                argv = text
            else:
                try:
                    argv = shlex.split(text)
                except ValueError as exc:
                    raise NodeError(f"command 解析失败: {exc}") from exc

        cwd = str(self.param("cwd", "") or "").strip()
        if cwd and not Path(cwd).is_dir():
            raise NodeError(f"工作目录不存在: {cwd}")

        timeout = float(self.param("timeout", 0) or 0) or DEFAULT_SHELL_TIMEOUT
        env_keys = [str(k) for k in (self.param("env", {}) or {}).keys()]
        env = dict(self.resolve("__env__", {}) or {}) if env_keys else None

        started = time.monotonic()
        try:
            proc = subprocess.run(  # noqa: S603
                argv,
                shell=use_shell,
                cwd=cwd or None,
                env=env,
                capture_output=True,
                text=True,
                timeout=timeout,
            )
        except FileNotFoundError as exc:
            raise NodeError(f"命令不存在: {argv[0]!r}") from exc
        except subprocess.TimeoutExpired as exc:
            raise NodeError(f"命令超时（>{timeout}s）: {' '.join(argv)[:120]}") from exc

        duration_ms = int((time.monotonic() - started) * 1000)
        stdout = proc.stdout or ""
        stderr = proc.stderr or ""
        max_chars = int(self.param("max_output_chars", 20000) or 20000)

        return {
            "exit_code": proc.returncode,
            "stdout": stdout[:max_chars],
            "stderr": stderr[:max_chars],
            "duration_ms": duration_ms,
            "command": argv if isinstance(argv, str) else " ".join(argv),
            # The executor routes on a node's declared port, so a non-zero exit
            # is a branch rather than a failure: a script that exits 1 to mean
            # "not found yet" is a normal thing to write.
            "__branch__": "ok" if proc.returncode == 0 else "error",
        }


class ExcelNode(BaseNode):
    """Read a cell or a range from an .xlsx workbook, or write one.

    Requires openpyxl. This node deliberately does not touch the WeChat
    automation state machine, so a spreadsheet step can be tested on its own.
    """

    def execute(self) -> dict[str, Any]:
        try:
            from openpyxl import load_workbook  # noqa: PLC0415
        except ImportError as exc:  # pragma: no cover - depends on install
            raise NodeError(
                "缺少 openpyxl，无法使用 excel 节点。请在项目虚拟环境执行: "
                "uv pip install --python .venv/bin/python openpyxl"
            ) from exc

        path = Path(str(self.param("path", "") or "")).expanduser()
        if not path.is_file():
            raise NodeError(f"Excel 文件不存在: {path}")

        sheet_name = str(self.param("sheet", "") or "")
        cell = str(self.param("cell", "") or "").strip()
        mode = str(self.param("mode", "read") or "read")

        if mode == "write":
            return self._write(load_workbook, path, sheet_name, cell)

        workbook = load_workbook(path, data_only=True, read_only=True)
        try:
            sheet = self._pick_sheet(workbook, sheet_name)
            if cell:
                value = sheet[cell].value
                return {"value": value, "cell": cell, "sheet": sheet.title, "path": str(path)}
            rows = [
                [c.value for c in row]
                for row in sheet.iter_rows(
                    min_row=int(self.param("min_row", 1) or 1),
                    max_row=int(self.param("max_row", 0) or 0) or sheet.max_row,
                    min_col=int(self.param("min_col", 1) or 1),
                    max_col=int(self.param("max_col", 0) or 0) or sheet.max_column,
                )
            ]
            return {
                "rows": rows,
                "row_count": len(rows),
                "sheet": sheet.title,
                "sheets": list(workbook.sheetnames),
                "path": str(path),
            }
        finally:
            workbook.close()

    def _write(self, load_workbook: Any, path: Path, sheet_name: str, cell: str) -> dict[str, Any]:
        import json  # noqa: PLC0415

        if not cell:
            raise NodeError("写入模式需要 cell 参数，例如 B2")
        raw = self.param("value", None)
        if isinstance(raw, str):
            stripped = raw.strip()
            if stripped.startswith(("[", "{")):
                try:
                    raw = json.loads(stripped)
                except json.JSONDecodeError:
                    pass  # a literal string that happens to start with a brace

        workbook = load_workbook(path)
        try:
            sheet = self._pick_sheet(workbook, sheet_name, create=True)
            sheet[cell] = raw
            workbook.save(path)
        finally:
            workbook.close()
        return {"written": True, "cell": cell, "value": raw, "path": str(path)}

    def _pick_sheet(self, workbook: Any, name: str, create: bool = False) -> Any:
        if not name:
            return workbook.worksheets[0]
        if name in workbook.sheetnames:
            return workbook[name]
        if create:
            return workbook.create_sheet(title=name)
        raise NodeError(f"工作表 {name!r} 不存在，可用: {workbook.sheetnames}")


def register_all(registry: NodeRegistry) -> NodeRegistry:
    """Register every node type in this module. Idempotent."""
    add = registry.register
    add(
NodeSpec(
            type="run_shell",
            label="执行命令",
            category="系统",
            handler=RunShellNode,
            doc="执行一个外部程序并捕获输出。默认不经过 shell，参数用列表或 shlex 解析。",
            params=[
                ParamSpec("command", "variable", "命令", required=True,
                          help="参数列表（推荐），或一个字符串（用 shlex 解析，不走 shell）"),
                ParamSpec("cwd", "text", "工作目录", default=""),
                ParamSpec("shell", "bool", "经过 shell", default=False,
                          help="仅在需要管道/通配符时开启。开启后命令里的文件名可被执行"),
                ParamSpec("timeout", "number", "超时秒数", default=0,
                          help="0 表示用默认 300 秒"),
                ParamSpec("max_output_chars", "number", "输出截断", default=20000),
            ],
            outputs=["ok", "error"],
        )
    )
    add(
NodeSpec(
            type="excel",
            label="Excel 读写",
            category="系统",
            handler=ExcelNode,
            doc="读取单元格/区域或写入单元格。需要 openpyxl。",
            params=[
                ParamSpec("path", "text", "文件路径", required=True),
                ParamSpec("mode", "select", "模式", default="read",
                          choices=["read", "write"]),
                ParamSpec("sheet", "text", "工作表", default="",
                          help="留空用第一个工作表"),
                ParamSpec("cell", "text", "单元格", default="",
                          help="读取时留空则读整表；写入时必填，如 B2"),
                ParamSpec("value", "variable", "写入值", default=None,
                          help="写模式的值。JSON 文本会按结构解析"),
                ParamSpec("min_row", "number", "起始行", default=1),
                ParamSpec("max_row", "number", "结束行", default=0,
                          help="0 表示到表尾"),
                ParamSpec("min_col", "number", "起始列", default=1),
                ParamSpec("max_col", "number", "结束列", default=0,
                          help="0 表示到表尾"),
            ],
            outputs=["ok"],
        )
    )
    return registry
